import json
import sys
from openpyxl import load_workbook

def main(excel_file):
    # Load the Excel workbook
    workbook = load_workbook(filename=excel_file)
    sheet = workbook.active

    # Initialize an empty list to store the combined data
    combined_data = []

    # Iterate over each row in the sheet
    for row in sheet.iter_rows(min_row=2, values_only=True):  # Skip header row
        ksa3_data = row[2]  # Assuming 'KSA3' is in column C (index 2)
        if ksa3_data:
            # Parse the JSON data in 'KSA3' column
            ksa3_list = json.loads(ksa3_data)
            combined_data.extend(ksa3_list)

    # Generate the output JSON file name
    output_file = excel_file.replace('.xlsx', '.json')

    # Save the combined data to a JSON file
    with open(output_file, 'w') as json_file:
        json.dump(combined_data, json_file)

    print(f"Combined data saved to {output_file}")

if __name__ == '__main__':
    if len(sys.argv) != 2:
        print("Usage: python gpt_loop_4_combine.py <excel_file>")
        sys.exit(1)

    excel_file = sys.argv[1]
    main(excel_file)
